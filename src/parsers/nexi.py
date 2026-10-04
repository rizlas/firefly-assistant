import warnings

import openpyxl

from src.firefly.models import Transaction

from .base import BaseParser
from .types import TransactionType, TransferType

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


class NexiParser(BaseParser):
    def parse(self, file_path, skip_already_imported=True):
        """
        Parse the Excel file of transactions.

        Args:
            file_path: Path to the Excel file
            skip_already_imported: If True, skips transactions already imported
        """
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        card_info = ws["B8"].value

        skip_states = {s.strip().lower() for s in self.behavior.skip_states}

        transactions = []
        for row in ws.iter_rows(min_row=11):
            date = row[2].value  # Column C
            reference = row[3].value  # Column D
            raw_account = row[5].value  # Column F
            state = row[6].value  # Column G
            total = row[9].value  # Column J

            # Rows not yet settled carry a provisional date and amount and have
            # no reference: skipping them means they get imported on a later
            # run, with their final values and a stable external_id.
            if (state or "").strip().lower() in skip_states:
                continue

            # Determine the transaction type, normalize the amount, and map the account
            tx_type, transfer_type, total = self.classify_transaction(
                total, raw_account
            )

            asset_account = self.asset_account
            mapped_account = self.firefly_client.map_account(
                self.firefly_client.ASSET_ACCOUNT_TYPE_MAP[tx_type], raw_account
            )

            # Use mapped account if found, otherwise fallback to raw account name
            destination_account = mapped_account or raw_account

            # For refund and transfer, source and destination must be swapped
            if tx_type in [
                TransactionType.REFUND,
                TransactionType.TRANSFER,
            ]:
                # Swap if it's a refund or the transfer type is IN
                if (
                    tx_type is TransactionType.REFUND
                    or transfer_type == TransferType.IN
                ):
                    asset_account = destination_account
                    destination_account = self.asset_account

            category = self.auto_categorize(destination_account)

            # The bank reference is stable across exports; fall back to the
            # computed hash for rows that have none (e.g. branch top-ups).
            tx = Transaction(
                id=reference,
                date=date,
                description="",
                source_account=asset_account,
                dest_account=destination_account,
                account_mapping=raw_account,
                total=total,
                type=tx_type.value,
                transfer_type=transfer_type.value if transfer_type else None,
                state=state,
                card=card_info,
                category=category,
            )

            if skip_already_imported:
                if self.firefly_client.transaction_exists(tx.id, tx.generate_id()):
                    continue

            transactions.append(tx)

        return transactions
